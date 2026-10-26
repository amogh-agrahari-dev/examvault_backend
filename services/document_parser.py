import os
import uuid
import logging
import asyncio
import pymupdf
import docx
import pptx
from typing import Tuple, List
from services.vision_service import (
    is_important_image,
    save_extracted_image,
    analyze_image_with_vision,
    transcribe_scanned_paper_image
)

logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = {
    '.txt', '.text', '.md', '.markdown', '.rst',
    '.py', '.java', '.cpp', '.c', '.h', '.hpp', '.cs',
    '.js', '.jsx', '.ts', '.tsx', '.html', '.htm', '.css',
    '.json', '.csv', '.tsv', '.sql', '.xml', '.yaml', '.yml',
    '.sh', '.bash', '.zsh', '.env', '.log'
}

IMAGE_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.webp', '.bmp', '.tiff'}


async def extract_document_content_and_images(
    file_path: str,
    mime_type: str = "",
    doc_id: uuid.UUID = None
) -> Tuple[str, List[str]]:
    """
    Extracts text AND embedded images from documents.
    Important diagrams/charts/figures are saved to disk, transcribed via vision AI,
    and embedded directly into the text stream and chunks so LLM and RAG understand them.
    """
    doc_id = doc_id or uuid.uuid4()
    ext = os.path.splitext(file_path)[1].lower().strip()
    mime = (mime_type or "").lower().split(';')[0].strip()

    try:
        if ext == '.pdf' or mime == 'application/pdf':
            return await _extract_from_pdf_with_images(file_path, doc_id)
        elif ext in {'.docx', '.doc'} or 'wordprocessingml' in mime or 'msword' in mime:
            return await _extract_from_docx_with_images(file_path, doc_id)
        elif ext in {'.pptx', '.ppt'} or 'presentationml' in mime or 'powerpoint' in mime:
            return await _extract_from_pptx_with_images(file_path, doc_id)
        elif ext in IMAGE_EXTENSIONS or mime.startswith('image/'):
            return await _extract_from_image_with_vision(file_path, doc_id)
        elif ext in TEXT_EXTENSIONS or mime.startswith('text/') or 'json' in mime or 'xml' in mime or 'csv' in mime:
            text = _extract_from_txt(file_path)
            return text, []
        else:
            try:
                text = _extract_from_txt(file_path)
                if text and len(text.strip()) > 0:
                    return text, []
            except Exception:
                pass
            raise ValueError(f"Unsupported file format: {ext or mime_type}. Supported: PDF, DOCX, PPTX, TXT, MD, Code, Images.")
    except Exception as e:
        logger.error(f"Error extracting content and images from {file_path}: {e}")
        raise


def extract_text(file_path: str, mime_type: str = "") -> str:
    """Synchronous fallback for text-only extraction."""
    ext = os.path.splitext(file_path)[1].lower().strip()
    mime = (mime_type or "").lower().split(';')[0].strip()

    if ext == '.pdf' or mime == 'application/pdf':
        return _extract_from_pdf_sync(file_path)
    elif ext in {'.docx', '.doc'} or 'wordprocessingml' in mime:
        return _extract_from_docx_sync(file_path)
    elif ext in {'.pptx', '.ppt'} or 'presentationml' in mime:
        return _extract_from_pptx_sync(file_path)
    elif ext in TEXT_EXTENSIONS or mime.startswith('text/'):
        return _extract_from_txt(file_path)
    elif ext in IMAGE_EXTENSIONS or mime.startswith('image/'):
        return f"Image Study Material: {os.path.basename(file_path)}"
    else:
        return _extract_from_txt(file_path)


def chunk_text(text: str, max_words: int = 1000) -> list[str]:
    """Splits a large string of text into smaller chunks for semantic LLM processing."""
    if not text or not text.strip():
        return ["No content could be extracted from this document."]

    words = text.split()
    if not words:
        return ["No content could be extracted from this document."]

    chunks = []
    for i in range(0, len(words), max_words):
        chunk = " ".join(words[i:i + max_words]).strip()
        if chunk:
            chunks.append(chunk)

    return chunks if chunks else [text.strip()[:1000]]


async def _extract_from_pdf_with_images(file_path: str, doc_id: uuid.UUID) -> Tuple[str, List[str]]:
    """Extracts text, formulas, and embedded figures from PDF pages."""
    full_text = ""
    extracted_image_urls = []
    max_images_to_analyze = 8

    try:
        doc = pymupdf.open(file_path)
        for page_num, page in enumerate(doc, start=1):
            page_text = page.get_text() or ""
            page_image_markdowns = []

            # 1. Extract embedded figures on page
            image_list = page.get_images(full=True)
            for img_info in image_list:
                if len(extracted_image_urls) >= max_images_to_analyze:
                    break

                xref = img_info[0]
                try:
                    base_image = doc.extract_image(xref)
                    image_bytes = base_image.get("image")
                    ext = base_image.get("ext", "png")

                    if is_important_image(image_bytes, min_dim=130, min_bytes=3500):
                        idx = len(extracted_image_urls) + 1
                        _, public_url = save_extracted_image(doc_id, image_bytes, index=idx, ext=ext)
                        extracted_image_urls.append(public_url)

                        # Analyze visual diagram via vision LLM
                        context = f"Page {page_num} of {os.path.basename(file_path)}"
                        vision_desc = await analyze_image_with_vision(image_bytes, context_hint=context)
                        if vision_desc:
                            page_image_markdowns.append(
                                f"\n\n[Embedded Figure / Diagram: {public_url}]\n"
                                f"Visual Diagram Explanation:\n{vision_desc}\n"
                            )
                        else:
                            page_image_markdowns.append(
                                f"\n\n[Embedded Figure / Diagram: {public_url}]\n(Visual figure from Page {page_num})\n"
                            )
                except Exception as img_err:
                    logger.debug(f"Skipping PDF image xref {xref}: {img_err}")

            # 2. Scanned / image-only page fallback
            if not page_text.strip() and not page_image_markdowns and len(extracted_image_urls) < max_images_to_analyze:
                try:
                    pix = page.get_pixmap(dpi=150)
                    pix_bytes = pix.tobytes("png")
                    if is_important_image(pix_bytes, min_dim=150, min_bytes=4000):
                        idx = len(extracted_image_urls) + 1
                        _, public_url = save_extracted_image(doc_id, pix_bytes, index=idx, ext="png")
                        extracted_image_urls.append(public_url)

                        vision_desc = await analyze_image_with_vision(pix_bytes, context_hint=f"Scanned Page {page_num}")
                        page_image_markdowns.append(
                            f"\n\n[Embedded Figure / Diagram: {public_url}]\n"
                            f"Scanned Page Content:\n{vision_desc or 'Scanned notes page'}\n"
                        )
                except Exception as scan_err:
                    logger.debug(f"Scanned page extraction failed: {scan_err}")

            # Combine page text with embedded figures
            combined_page = page_text
            if page_image_markdowns:
                combined_page += "\n" + "\n".join(page_image_markdowns)

            if combined_page.strip():
                full_text += combined_page + "\n\n"

        doc.close()
    except Exception as e:
        logger.error(f"Error reading PDF {file_path}: {e}")
        raise

    if not full_text.strip():
        full_text = f"Document: {os.path.basename(file_path)}. Contains visual or unextractable text."

    return full_text.strip(), extracted_image_urls


async def _extract_from_docx_with_images(file_path: str, doc_id: uuid.UUID) -> Tuple[str, List[str]]:
    """Extracts text, tables, and embedded images from DOCX."""
    full_text = ""
    extracted_image_urls = []
    max_images_to_analyze = 6

    try:
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            if para.text.strip():
                full_text += para.text + "\n"

        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text += " | ".join(row_text) + "\n"

        # Extract embedded images from document relationships
        for rel in doc.part.rels.values():
            if len(extracted_image_urls) >= max_images_to_analyze:
                break
            if "image" in rel.target_ref:
                try:
                    img_part = rel.target_part
                    img_bytes = img_part.blob
                    ext = os.path.splitext(rel.target_ref)[1].replace(".", "") or "png"

                    if is_important_image(img_bytes, min_dim=120, min_bytes=3500):
                        idx = len(extracted_image_urls) + 1
                        _, public_url = save_extracted_image(doc_id, img_bytes, index=idx, ext=ext)
                        extracted_image_urls.append(public_url)

                        vision_desc = await analyze_image_with_vision(img_bytes, context_hint=os.path.basename(file_path))
                        if vision_desc:
                            full_text += f"\n\n[Embedded Figure / Diagram: {public_url}]\nVisual Content:\n{vision_desc}\n\n"
                        else:
                            full_text += f"\n\n[Embedded Figure / Diagram: {public_url}]\n(Important visual figure)\n\n"
                except Exception as rel_err:
                    logger.debug(f"Skipping DOCX image part: {rel_err}")

    except Exception as e:
        logger.error(f"Error reading DOCX {file_path}: {e}")
        raise

    return full_text.strip() or f"Document: {os.path.basename(file_path)}", extracted_image_urls


async def _extract_from_pptx_with_images(file_path: str, doc_id: uuid.UUID) -> Tuple[str, List[str]]:
    """Extracts text and picture shapes from presentation slides."""
    full_text = ""
    extracted_image_urls = []
    max_images_to_analyze = 8

    try:
        prs = pptx.Presentation(file_path)
        for slide_num, slide in enumerate(prs.slides, start=1):
            slide_content = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_content.append(shape.text.strip())

                # Extract pictures from slides
                if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE and len(extracted_image_urls) < max_images_to_analyze:
                    try:
                        img_bytes = shape.image.blob
                        ext = shape.image.ext or "png"
                        if is_important_image(img_bytes, min_dim=120, min_bytes=3500):
                            idx = len(extracted_image_urls) + 1
                            _, public_url = save_extracted_image(doc_id, img_bytes, index=idx, ext=ext)
                            extracted_image_urls.append(public_url)

                            vision_desc = await analyze_image_with_vision(img_bytes, context_hint=f"Slide {slide_num}")
                            if vision_desc:
                                slide_content.append(f"[Embedded Figure / Diagram: {public_url}]\n{vision_desc}")
                            else:
                                slide_content.append(f"[Embedded Figure / Diagram: {public_url}]\n(Figure on slide {slide_num})")
                    except Exception as shape_err:
                        logger.debug(f"Skipping slide shape picture: {shape_err}")

            if slide_content:
                full_text += f"Slide {slide_num}:\n" + "\n".join(slide_content) + "\n\n"

    except Exception as e:
        logger.error(f"Error reading PPTX {file_path}: {e}")
        raise

    return full_text.strip() or f"Presentation: {os.path.basename(file_path)}", extracted_image_urls


async def _extract_from_image_with_vision(file_path: str, doc_id: uuid.UUID) -> Tuple[str, List[str]]:
    """Analyzes an uploaded standalone image file with vision AI."""
    try:
        with open(file_path, "rb") as f:
            img_bytes = f.read()

        ext = os.path.splitext(file_path)[1].replace(".", "") or "png"
        _, public_url = save_extracted_image(doc_id, img_bytes, index=1, ext=ext)

        vision_desc = await analyze_image_with_vision(img_bytes, context_hint=os.path.basename(file_path))
        full_text = (
            f"[Embedded Figure / Diagram: {public_url}]\n\n"
            f"Visual Diagram & Content:\n{vision_desc or 'Uploaded study diagram'}\n\n"
        )
        return full_text.strip(), [public_url]
    except Exception as e:
        logger.error(f"Error analyzing image file {file_path}: {e}")
        return f"Image Document: {os.path.basename(file_path)}", []


def _extract_from_pdf_sync(file_path: str) -> str:
    text = ""
    try:
        doc = pymupdf.open(file_path)
        for page in doc:
            p_text = page.get_text() or ""
            if p_text.strip():
                text += p_text + "\n\n"
        doc.close()
    except Exception:
        pass
    return text.strip() or f"Document: {os.path.basename(file_path)}"


def _extract_from_docx_sync(file_path: str) -> str:
    text = ""
    try:
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            if para.text.strip():
                text += para.text + "\n"
    except Exception:
        pass
    return text.strip() or f"Document: {os.path.basename(file_path)}"


def _extract_from_pptx_sync(file_path: str) -> str:
    text = ""
    try:
        prs = pptx.Presentation(file_path)
        for s in prs.slides:
            for shp in s.shapes:
                if hasattr(shp, "text") and shp.text.strip():
                    text += shp.text.strip() + "\n"
    except Exception:
        pass
    return text.strip() or f"Presentation: {os.path.basename(file_path)}"


def _extract_from_txt(file_path: str) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            with open(file_path, "r", encoding=encoding) as f:
                content = f.read().strip()
                if content:
                    return content
        except (UnicodeDecodeError, Exception):
            continue
    return f"Text Document: {os.path.basename(file_path)}"


async def extract_past_paper_content(file_path: str, mime_type: str = "") -> str:
    """
    Extracts text from previous year question papers.
    Seamlessly handles:
    1. Digital PDFs (fast text extraction).
    2. Scanned PDFs / Hardcopy photos (automatic vision OCR on each page).
    3. Direct image uploads (PNG, JPG, JPEG, WEBP exam photos).
    4. DOCX and TXT papers.
    """
    ext = os.path.splitext(file_path)[1].lower().strip()
    mime = (mime_type or "").lower().split(';')[0].strip()

    # Case A: Standalone Image (Photograph / Scan of paper)
    if ext in IMAGE_EXTENSIONS or mime.startswith('image/'):
        try:
            with open(file_path, "rb") as f:
                img_bytes = f.read()
            transcribed = await transcribe_scanned_paper_image(
                img_bytes, page_hint=f"Physical Question Paper: {os.path.basename(file_path)}"
            )
            return transcribed if transcribed else f"Scanned Examination Paper Image: {os.path.basename(file_path)}"
        except Exception as e:
            logger.error(f"Error reading image past paper {file_path}: {e}")
            return f"Scanned Examination Paper: {os.path.basename(file_path)}"

    # Case B: PDF Document (Digital or Scanned hardcopy)
    if ext == '.pdf' or mime == 'application/pdf':
        digital_text = ""
        scanned_pages = []
        try:
            doc = pymupdf.open(file_path)
            total_pages = len(doc)
            for page_num, page in enumerate(doc, start=1):
                p_text = page.get_text() or ""
                # If page has substantial selectable text, collect it
                if len(p_text.strip()) > 50:
                    digital_text += f"\n--- Page {page_num} ---\n" + p_text
                else:
                    # Scanned or image-only page: render to image for vision OCR
                    try:
                        pix = page.get_pixmap(dpi=150)
                        scanned_pages.append((page_num, pix.tobytes("jpeg")))
                    except Exception as pix_err:
                        logger.warning(f"Failed to render page {page_num} to pixmap: {pix_err}")
            doc.close()
        except Exception as pdf_err:
            logger.error(f"Failed to read PDF {file_path}: {pdf_err}")

        # If purely digital PDF with full selectable text and no scanned pages
        if len(digital_text.strip()) > 100 and not scanned_pages:
            return digital_text.strip()

        # If scanned pages exist (scanned copy of hardcopy)
        if scanned_pages:
            logger.info(
                f"Detected {len(scanned_pages)} scanned hardcopy page(s) in {os.path.basename(file_path)}. "
                "Activating vision OCR..."
            )
            ocr_text = ""
            for page_num, img_bytes in scanned_pages:
                page_transcription = await transcribe_scanned_paper_image(
                    img_bytes,
                    page_hint=f"Page {page_num} of {total_pages} of scanned physical examination paper"
                )
                if page_transcription:
                    ocr_text += f"\n--- Page {page_num} (Transcribed from Scanned Hardcopy) ---\n" + page_transcription + "\n"

            combined = (digital_text + "\n" + ocr_text).strip()
            if combined:
                return combined

        return digital_text.strip() if digital_text.strip() else f"Scanned Examination Paper: {os.path.basename(file_path)}"

    # Case C: Fallback to synchronous text parser (DOCX, TXT, etc.)
    return extract_text(file_path, mime_type)
